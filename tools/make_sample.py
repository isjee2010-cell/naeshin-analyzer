# 가상 나이스 교과학습발달상황 파일 생성 (테스트용)
# 이름은 흔한 성 + 이름 글자를 무작위로 조합한 가상 인물이며, 성적도 난수로 만든 것입니다.
# 사용법: python3 tools/make_sample.py <출력폴더> g2|g3|g1|rand <시드>|xss
import openpyxl, random, sys
from openpyxl.styles import Font
random.seed(int(sys.argv[3]) if len(sys.argv)>3 else 7)
SURN="김이박최정강조윤장임한오서신권황안송류홍"
GIV=["민준","서연","도윤","하은","지호","수아","예준","지유","시우","서윤","주원","하린","유준","채원","지훈","다은","건우","소율","우진","지민","현우","예린","선우","수빈","연우","가은","정우","나연"]
def mk(ban_hak, ban_no, nstu, scale, terms, path, fmt_variant=0, skipp=0.0):
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="Sheet1"
    ws['A1']="교과학습발달상황"; ws['A2']=f"2026학년도 {ban_hak}학년 {ban_no}반"
    hdr=["번 호","성명","학년","학기","교과","과목","학점","원점수/과목평균","성취도","성취도별분포비율","석차등급","수강자수","비고"]
    r=4
    for j,h in enumerate(hdr): ws.cell(r,j+1,h)
    r+=1
    subj={ (1,1):[("국어","공통국어1",4),("수학","공통수학1",4),("영어","공통영어1",4),("사회(역사/도덕포함)","한국사1",3),("사회(역사/도덕포함)","통합사회1",4),("과학","통합과학1",4),("과학","과학탐구실험1",1),("기술·가정/정보","정보",2)],
           (1,2):[("국어","공통국어2",4),("수학","공통수학2",4),("영어","공통영어2",4),("사회(역사/도덕포함)","한국사2",3),("사회(역사/도덕포함)","통합사회2",4),("과학","통합과학2",4),("과학","과학탐구실험2",1),("제2외국어/한문","한문",2)],
           (2,1):[("국어","문학",4),("수학","대수",4),("영어","영어Ⅰ",4),("사회(역사/도덕포함)","세계시민과 지리",3),("과학","물리학",3),("과학","화학",3),("과학","생명과학",3),("기술·가정/정보","인공지능 기초",3)],
           (2,2):[("국어","독서와 작문",4),("수학","미적분Ⅰ",4),("영어","영어Ⅱ",4),("사회(역사/도덕포함)","사회와 문화",3),("과학","지구과학",3)],
           (3,1):[("국어","화법과 언어",4),("수학","확률과 통계",4),("영어","영어 독해와 작문",4),("과학","물리학Ⅱ",2)]}
    if scale==9:
        subj={ (1,1):[("국어","국어",4),("수학","수학",4),("영어","영어",4),("사회(역사/도덕포함)","한국사",3),("사회(역사/도덕포함)","통합사회",4),("과학","통합과학",4)],
               (1,2):[("국어","국어",4),("수학","수학",4),("영어","영어",4),("사회(역사/도덕포함)","한국사",3),("사회(역사/도덕포함)","통합사회",4),("과학","통합과학",4)],
               (2,1):[("국어","문학",4),("수학","수학Ⅰ",4),("영어","영어Ⅰ",4),("사회(역사/도덕포함)","생활과 윤리",3),("과학","화학Ⅰ",3)],
               (2,2):[("국어","독서",4),("수학","수학Ⅱ",4),("영어","영어Ⅱ",4),("사회(역사/도덕포함)","윤리와 사상",3),("과학","생명과학Ⅰ",3)],
               (3,1):[("국어","언어와 매체",4),("수학","미적분",4),("영어","영어 독해와 작문",4),("과학","화학Ⅱ",2)]}
    studs=[]
    for i in range(1,nstu+1):
        nm=random.choice(SURN)+random.choice(GIV)
        base=random.uniform(1.0, scale-0.3); trend=random.uniform(-0.5,0.5)*(2 if scale==9 else 1)
        weak=random.choice(["국어","수학","영어","과학",None,None,None])
        studs.append((i,nm,base,trend,weak))
    avgs={}
    for (h,s),lst in subj.items():
        for g,sb,cr in lst: avgs[(h,s,sb)]=round(random.uniform(55,78),1)
    for (no,nm,base,trend,weak) in studs:
        first=True
        skip0 = random.random()<skipp
        for ti,(h,s) in enumerate(terms):
            if skip0 and ti==0: continue
            firstTerm=True
            for g,sb,cr in subj[(h,s)]:
                lvl=base+trend*ti+random.uniform(-0.6,0.6)+(1.6 if weak and weak in g else 0)
                rk=max(1,min(scale,round(lvl)))
                av=avgs[(h,s,sb)]; sc=max(10,min(100,round(av+(scale/2-rk)*(36/scale)+random.uniform(-5,5))))
                ach="ABCDE"[min(4,max(0,int((rk-1)*5/scale)))]
                row=[no if first else None, nm if first else None, h if firstTerm else None, s if firstTerm else None, g, sb, cr, f"{sc}/{av}", ach, "", rk, 120+random.randint(-20,20), ""]
                if "탐구실험" in sb: row[10]=""; row[7]=""; row[11]=""; row[8]="P" if scale==5 else "A"
                if sb=="물리학Ⅱ": row[11]=14
                for j,v in enumerate(row):
                    if v is not None: ws.cell(r,j+1,v)
                r+=1; first=False; firstTerm=False
            ws.cell(r,1,"이수학점 합계"); ws.cell(r,7,sum(c for _,_,c in subj[(h,s)])); r+=1
    # 체육·예술
    r+=1; ws.cell(r,1,"< 체육ㆍ예술 >" if fmt_variant else "체육·예술"); r+=1
    for j,h in enumerate(["번 호","성명","학년","학기","교과","과목","학점","성취도"]): ws.cell(r,j+1,h)
    r+=1
    for (no,nm,*_ ) in studs:
        first=True
        for (h,s) in terms:
            for g,sb in [("체육","체육" if h==1 else "운동과 건강"),("예술","음악" if s==1 else "미술")]:
                row=[no if first else None, nm if first else None,h,s,g,sb,2,random.choice("AAB")]
                for j,v in enumerate(row):
                    if v is not None: ws.cell(r,j+1,v)
                r+=1; first=False
    r+=1; ws.cell(r,1,"교양교과"); r+=1
    for j,h in enumerate(["번 호","성명","학년","학기","교과","과목","학점","이수여부"]): ws.cell(r,j+1,h)
    r+=1
    for (no,nm,*_ ) in studs:
        row=[no,nm,terms[-1][0],terms[-1][1],"교양","진로와 직업",2,"P"]
        for j,v in enumerate(row): ws.cell(r,j+1,v)
        r+=1
    wb.save(path)
if __name__=="__main__":
    out=sys.argv[1]; kind=sys.argv[2]
    import os; os.makedirs(out, exist_ok=True)
    if kind=="g2":
        for b,n in [(1,24),(2,23),(3,25)]:
            random.seed(100+b); mk(2,b,n,5,[(1,1),(1,2),(2,1)],f"{out}/neis_2_{b}.xlsx")
    elif kind=="g3":
        random.seed(300); mk(3,4,20,9,[(1,1),(1,2),(2,1),(2,2),(3,1)],f"{out}/neis_3_4.xlsx")
    elif kind=="rand":
        sd=int(sys.argv[3])
        for b in (1,2,3):
            random.seed(sd*10+b); mk(2,b,30,5,[(1,1),(1,2),(2,1)],f"{out}/r_{sd}_{b}.xlsx",skipp=0.08)
    elif kind=="xss":
        # 보안 시험용: 이름·과목·교과 칸에 공격 문자열을 넣은 파일
        random.seed(101); path=f"{out}/neis_xss.xlsx"; mk(2,1,24,5,[(1,1),(1,2),(2,1)],path)
        P=['<img src=x onerror="window.__xss=1">','"><svg onload=window.__xss=2>','</'+'script><script>window.__xss=3</'+'script>',
           '=HYPERLINK("http://evil.example","x")',"';window.__xss=5;//",'__proto__']
        wb=openpyxl.load_workbook(path); ws=wb.active; k=0; c=0
        for row in ws.iter_rows(min_row=5):
            if row[1].value and k<len(P): row[1].value=P[k]; k+=1
            if row[5].value=='공통국어1' and c<3: row[5].value=['<img src=x onerror=window.__xss=11>','constructor','{{7*7}}${7*7}'][c]; c+=1
            if row[4].value=='기술·가정/정보': row[4].value='__proto__'
            if row[5].value=='공통영어2': row[4].value='<svg onload=window.__xss=12>'
        wb.save(path)
    elif kind=="g1":
        random.seed(200); mk(1,2,22,5,[(1,1),(1,2)],f"{out}/neis_1_2.xlsx")
